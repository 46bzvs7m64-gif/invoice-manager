"""发票管理接口

阶段一：列表 / 详情 / 下载 / 预览 / 删除
阶段二：手动编辑（PUT）
阶段三：多条件筛选排序、标记使用（单条/批量）
"""
import hashlib
import re
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Optional
from urllib.parse import quote

from fastapi import APIRouter, Depends, File, HTTPException, Response, UploadFile
from fastapi.responses import FileResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from sqlalchemy.orm import Session

from ..config import settings
from ..database import get_db
from ..models import Invoice
from ..schemas import BatchUse, InvoiceUpdate, UseIn

# 手动上传文件大小上限（20MB）
MAX_UPLOAD_SIZE = 20 * 1024 * 1024

router = APIRouter()


def _to_dict(inv: Invoice) -> dict:
    return {
        "id": inv.id,
        "invoice_no": inv.invoice_no,
        "invoice_type": inv.invoice_type,
        "invoice_date": inv.invoice_date,
        "amount_total": inv.amount_total,
        "amount_without_tax": inv.amount_without_tax,
        "tax_amount": inv.tax_amount,
        "seller_name": inv.seller_name,
        "seller_tax_no": inv.seller_tax_no,
        "buyer_name": inv.buyer_name,
        "buyer_tax_no": inv.buyer_tax_no,
        "category": inv.category,
        "file_name": inv.file_name,
        "file_size": inv.file_size or 0,
        "source_email": inv.source_email,
        "status": inv.status,
        "used_remark": inv.used_remark,
        "used_at": inv.used_at.strftime("%Y-%m-%d %H:%M") if inv.used_at else None,
        "ocr_status": inv.ocr_status,
        "ocr_error": inv.ocr_error,
        "created_at": inv.created_at.strftime("%Y-%m-%d %H:%M") if inv.created_at else "",
    }


def _get_inv(db: Session, invoice_id: int) -> Invoice:
    inv = db.get(Invoice, invoice_id)
    if not inv:
        raise HTTPException(404, "发票不存在")
    return inv


def _file_response(inv: Invoice, media_type: str | None = None) -> FileResponse:
    p = Path(inv.file_path)
    if not p.exists():
        raise HTTPException(404, "发票文件不存在或已被移动")
    return FileResponse(str(p), filename=inv.file_name, media_type=media_type)


def _apply_filters(q, **f):
    """组合多条件筛选"""
    if f["status"] in ("unused", "used"):
        q = q.filter(Invoice.status == f["status"])
    if f["seller"]:
        q = q.filter(Invoice.seller_name == f["seller"].strip())
    if f["buyer"]:
        q = q.filter(Invoice.buyer_name == f["buyer"].strip())
    # invoice_date 以 YYYY-MM-DD 字符串存储，字典序与时间序一致，可直接比较
    if f["date_from"]:
        q = q.filter(Invoice.invoice_date >= f["date_from"])
    if f["date_to"]:
        q = q.filter(Invoice.invoice_date <= f["date_to"])
    if f["amount_min"] is not None:
        q = q.filter(Invoice.amount_total >= f["amount_min"])
    if f["amount_max"] is not None:
        q = q.filter(Invoice.amount_total <= f["amount_max"])
    if f["invoice_type"]:
        q = q.filter(Invoice.invoice_type.like(f"%{f['invoice_type'].strip()}%"))
    if f["category"]:
        q = q.filter(Invoice.category == f["category"].strip())
    if f["keyword"]:
        kw = f"%{f['keyword'].strip()}%"
        q = q.filter(
            Invoice.file_name.like(kw)
            | Invoice.invoice_no.like(kw)
            | Invoice.seller_name.like(kw)
            | Invoice.buyer_name.like(kw)
        )
    return q


@router.get("")
def list_invoices(
    status: str = "unused",
    seller: str = "",
    buyer: str = "",
    date_from: str = "",
    date_to: str = "",
    amount_min: Optional[float] = None,
    amount_max: Optional[float] = None,
    invoice_type: str = "",
    category: str = "",
    keyword: str = "",
    sort: str = "date",
    order: str = "desc",
    page: int = 1,
    page_size: int = 20,
    db: Session = Depends(get_db),
):
    """发票列表：多条件组合筛选 + 排序 + 分页；默认只显示「未使用」"""
    q = _apply_filters(
        db.query(Invoice),
        status=status, seller=seller, buyer=buyer,
        date_from=date_from, date_to=date_to,
        amount_min=amount_min, amount_max=amount_max,
        invoice_type=invoice_type, category=category, keyword=keyword,
    )
    # 排序：date=开票日期 / amount=价税合计 / id=采集顺序
    col = {"date": Invoice.invoice_date, "amount": Invoice.amount_total}.get(sort, Invoice.id)
    q = q.order_by(col.desc() if order == "desc" else col.asc())

    total = q.count()
    page = max(1, page)
    page_size = max(1, min(page_size, 100))
    rows = q.offset((page - 1) * page_size).limit(page_size).all()

    # 附带当前筛选结果的金额合计，方便核对（如报销总金额）
    matched_amount = round(
        sum(r.amount_total or 0 for r in _apply_filters(
            db.query(Invoice),
            status=status, seller=seller, buyer=buyer,
            date_from=date_from, date_to=date_to,
            amount_min=amount_min, amount_max=amount_max,
            invoice_type=invoice_type, category=category, keyword=keyword,
        ).all()),
        2,
    )
    return {
        "total": total, "page": page, "page_size": page_size,
        "matched_amount": matched_amount,
        "rows": [_to_dict(r) for r in rows],
    }


@router.get("/export")
def export_invoices(
    status: str = "unused",
    seller: str = "",
    buyer: str = "",
    date_from: str = "",
    date_to: str = "",
    amount_min: Optional[float] = None,
    amount_max: Optional[float] = None,
    invoice_type: str = "",
    category: str = "",
    keyword: str = "",
    mark_used: int = 0,
    db: Session = Depends(get_db),
):
    """按当前筛选条件导出 Excel 报销单；mark_used=1 时同时将导出的发票标记为已使用"""
    rows = (
        _apply_filters(
            db.query(Invoice),
            status=status, seller=seller, buyer=buyer,
            date_from=date_from, date_to=date_to,
            amount_min=amount_min, amount_max=amount_max,
            invoice_type=invoice_type, category=category, keyword=keyword,
        )
        .order_by(Invoice.invoice_date.asc(), Invoice.id.asc())
        .all()
    )
    if not rows:
        raise HTTPException(400, "当前筛选条件下没有可导出的发票")

    now = datetime.now()
    wb = Workbook()
    ws = wb.active
    ws.title = "报销单"
    headers = ["序号", "开票日期", "发票号码", "发票类型", "销售方", "购买方",
               "消费类目", "价税合计", "不含税金额", "税额", "使用状态", "使用备注", "使用时间"]
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.alignment = Alignment(horizontal="center")
    for i, inv in enumerate(rows, 1):
        ws.append([
            i,
            inv.invoice_date or "",
            inv.invoice_no or "",
            inv.invoice_type or "",
            inv.seller_name or "",
            inv.buyer_name or "",
            inv.category or "",
            inv.amount_total,
            inv.amount_without_tax,
            inv.tax_amount,
            "已使用" if inv.status == "used" else "未使用",
            inv.used_remark or "",
            inv.used_at.strftime("%Y-%m-%d %H:%M") if inv.used_at else "",
        ])
    total_amount = round(sum(r.amount_total or 0 for r in rows), 2)
    ws.append([])
    ws.append([f"合计：{len(rows)} 张发票，价税合计 ¥{total_amount}"])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    for col, w in zip("ABCDEFGHIJKLM", (6, 12, 24, 16, 32, 32, 10, 10, 12, 10, 10, 18, 16)):
        ws.column_dimensions[col].width = w

    buf = BytesIO()
    wb.save(buf)

    if mark_used:
        for inv in rows:
            if inv.status != "used":
                inv.status = "used"
                inv.used_remark = inv.used_remark or "导出报销单"
                inv.used_at = now
        db.commit()

    fname = f"报销单_{now:%Y%m%d_%H%M}.xlsx"
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(fname)}"},
    )


@router.post("/upload")
async def upload_invoice(file: UploadFile = File(...), db: Session = Depends(get_db)):
    """手动上传发票：校验格式/魔数/大小，MD5 去重后落盘入库（待识别状态）"""
    fname = file.filename or ""
    ext = Path(fname).suffix.lower()
    if ext not in (".pdf", ".ofd"):
        raise HTTPException(400, "仅支持 PDF / OFD 格式的发票文件")

    payload = await file.read()
    if not payload:
        raise HTTPException(400, "文件内容为空")
    if len(payload) > MAX_UPLOAD_SIZE:
        raise HTTPException(400, "文件过大，单文件不能超过 20MB")
    # 魔数校验，防止改名的伪造文件
    if ext == ".pdf" and not payload.startswith(b"%PDF-"):
        raise HTTPException(400, "文件不是有效的 PDF（内容校验失败）")
    if ext == ".ofd" and not payload.startswith(b"PK"):
        raise HTTPException(400, "文件不是有效的 OFD（内容校验失败）")

    # MD5 去重：与邮箱采集共用同一判重标准
    md5 = hashlib.md5(payload).hexdigest()
    if db.query(Invoice).filter(Invoice.file_md5 == md5).first():
        raise HTTPException(409, "该发票已存在，重复文件已跳过")

    now = datetime.now()
    safe_name = re.sub(r'[\\/:*?"<>|\r\n\t]', "_", fname).strip() or f"invoice{ext}"
    save_dir = Path(settings.STORAGE_PATH) / str(now.year) / "manual_upload"
    save_dir.mkdir(parents=True, exist_ok=True)
    file_path = save_dir / f"{now:%Y%m%d_%H%M%S}_{safe_name}"
    n = 1
    while file_path.exists():
        file_path = save_dir / f"{now:%Y%m%d_%H%M%S}_{n}_{safe_name}"
        n += 1
    file_path.write_bytes(payload)

    inv = Invoice(
        file_path=str(file_path),
        file_name=fname,
        file_size=len(payload),
        file_md5=md5,
        source_email="手动上传",
        status="unused",
        ocr_status="pending",
    )
    db.add(inv)
    db.commit()
    db.refresh(inv)
    return {"ok": True, "id": inv.id, "file_name": fname,
            "message": "上传成功，可到「智能识别」页识别发票内容"}


@router.get("/{invoice_id}")
def invoice_detail(invoice_id: int, db: Session = Depends(get_db)):
    inv = _get_inv(db, invoice_id)
    return _to_dict(inv)


@router.put("/{invoice_id}")
def update_invoice(invoice_id: int, body: InvoiceUpdate, db: Session = Depends(get_db)):
    """手动编辑发票字段；空字符串视为清空"""
    inv = _get_inv(db, invoice_id)
    data = body.model_dump(exclude_unset=True)
    for k, v in data.items():
        setattr(inv, k, v if v != "" else None)
    if data:
        inv.ocr_status = "success"
        inv.ocr_error = None
    db.commit()
    return {"ok": True, "message": "发票信息已更新"}


@router.delete("/{invoice_id}")
def delete_invoice(invoice_id: int, db: Session = Depends(get_db)):
    """删除发票记录并删除对应文件"""
    inv = _get_inv(db, invoice_id)
    try:
        p = Path(inv.file_path)
        if p.exists():
            p.unlink()
    except OSError:
        pass
    db.delete(inv)
    db.commit()
    return {"ok": True, "message": "已删除"}


@router.get("/{invoice_id}/file")
def download_file(invoice_id: int, db: Session = Depends(get_db)):
    inv = _get_inv(db, invoice_id)
    return _file_response(inv)


@router.get("/{invoice_id}/preview")
def preview_file(invoice_id: int, db: Session = Depends(get_db)):
    inv = _get_inv(db, invoice_id)
    media = "application/pdf" if inv.file_name.lower().endswith(".pdf") else "application/octet-stream"
    return _file_response(inv, media_type=media)


# ============ 阶段三：标记使用 ============

@router.post("/batch-use")
def batch_use(body: BatchUse, db: Session = Depends(get_db)):
    """批量标记已使用"""
    if not body.ids:
        raise HTTPException(400, "请选择至少一张发票")
    now = datetime.now()
    n = 0
    for i in body.ids:
        inv = db.get(Invoice, i)
        if inv and inv.status != "used":
            inv.status = "used"
            inv.used_remark = body.remark or None
            inv.used_at = now
            n += 1
    db.commit()
    return {"ok": True, "count": n, "message": f"已标记 {n} 张发票为已使用"}


@router.post("/{invoice_id}/use")
def mark_use(invoice_id: int, body: UseIn, db: Session = Depends(get_db)):
    """单张标记已使用，可填写使用备注（如报销单号）"""
    inv = _get_inv(db, invoice_id)
    inv.status = "used"
    inv.used_remark = body.remark or None
    inv.used_at = datetime.now()
    db.commit()
    return {"ok": True, "message": "已标记为已使用"}


@router.post("/{invoice_id}/unuse")
def mark_unuse(invoice_id: int, db: Session = Depends(get_db)):
    """取消使用标记，恢复为未使用"""
    inv = _get_inv(db, invoice_id)
    inv.status = "unused"
    inv.used_remark = None
    inv.used_at = None
    db.commit()
    return {"ok": True, "message": "已恢复为未使用"}
